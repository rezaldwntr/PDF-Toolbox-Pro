# app/routers/storage.py
"""
Router untuk Google Cloud Storage (GCS) Presigned Upload & Processing — Tier 3.

Endpoints:
  GET  /storage/status           — Cek apakah GCS aktif dan dapat digunakan
  POST /storage/presigned-upload — Minta Signed URL untuk direct upload dari browser
  POST /storage/process-job      — Perintahkan server memproses berkas yang sudah diunggah di GCS
"""

import os
import io
import re
import uuid
import logging
import asyncio
import tempfile
import multiprocessing
from typing import Optional, Dict, Any
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi import APIRouter, HTTPException, Header, status
from fastapi.responses import JSONResponse
from pydantic import BaseModel

import fitz  # PyMuPDF
from pdf2docx import Converter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
import pdfplumber
import pandas as pd
from openpyxl.styles import Border, Side, Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.core.config import MAX_FILE_SIZE_BY_TIER, GCS_BUCKET_NAME
from app.utils.file_utils import get_tier_limit, cleanup_folder, validate_pdf_bytes, get_safe_base_name
from app.utils.supabase_utils import verify_supabase_token
from app.utils.job_store import create_job, update_job
from app.utils.async_task import create_tracked_task
from app.utils.gcs_utils import (
    is_gcs_available,
    generate_upload_signed_url,
    download_blob_to_bytes,
    delete_blob,
    configure_bucket_security,
)
from app.routers.tools_helpers import translate_text_chunk, get_target_pages

router = APIRouter(prefix="/storage", tags=["Storage & Direct Upload"])


class PresignedRequest(BaseModel):
    filename: str
    file_size: int
    content_type: Optional[str] = "application/pdf"
    tier: Optional[str] = "free"


class ProcessJobRequest(BaseModel):
    blob_name: str
    action: str  # "word" | "excel" | "ppt" | "image" | "ocr" | "translate"
    options: Optional[Dict[str, Any]] = None


@router.get("/status")
def get_storage_status():
    """Mengecek apakah arsitektur GCS Direct Upload aktif dan siap digunakan."""
    available = is_gcs_available()
    return {
        "gcs_available": available,
        "bucket": GCS_BUCKET_NAME if available else None,
        "limits": {k: v // (1024 * 1024) for k, v in MAX_FILE_SIZE_BY_TIER.items()},
        "message": "GCS Direct Upload Aktif (Kapasitas hingga 500 MB)" if available else "Fallback ke Direct Multipart Upload (Maks 50-100 MB)",
    }


@router.post("/presigned-upload")
async def request_presigned_upload(
    req: PresignedRequest,
    authorization: Optional[str] = Header(None),
    x_user_tier: Optional[str] = Header(None)
):
    """
    Menghasilkan Signed URL aman agar klien dapat mengunggah berkas besar langsung ke GCS.
    Terproteksi validasi token kriptografis Supabase JWT untuk mencegah pemalsuan kuota tier.
    """
    # 1. Otentikasi Tier: Pengguna yang mengklaim tier Pro wajib menyertakan token Supabase valid
    effective_tier = "free"
    claimed_tier = (x_user_tier or req.tier or "free").lower().strip()

    if claimed_tier in ("flash", "monthly", "annual"):
        if authorization and authorization.strip().startswith("Bearer "):
            token = authorization.strip().split("Bearer ")[1].strip()
            verified_user = await verify_supabase_token(token)
            if verified_user:
                effective_tier = claimed_tier
            else:
                effective_tier = "free"
        else:
            effective_tier = "free"
    else:
        effective_tier = "free"

    tier_limit = get_tier_limit(effective_tier)
    if req.file_size <= 0:
        raise HTTPException(status_code=400, detail="Ukuran berkas tidak valid.")

    if req.file_size > tier_limit:
        limit_mb = tier_limit // (1024 * 1024)
        raise HTTPException(
            status_code=400,
            detail=f"Ukuran berkas ({req.file_size / (1024 * 1024):.1f} MB) melebihi batas paket Anda ({limit_mb} MB)."
        )

    # 2. Bersihkan dan sanitasi nama berkas (Path Traversal Protection)
    clean_name = re.sub(r'[^\w\-_\. ]', '_', req.filename).strip() or "dokumen.pdf"
    if not clean_name.lower().endswith(".pdf"):
        clean_name += ".pdf"

    unique_prefix = uuid.uuid4().hex[:12]
    blob_path = f"uploads/{unique_prefix}/{clean_name}"

    content_type = "application/pdf"

    # Jika GCS tidak tersedia di lokal/staging, beri tahu frontend untuk fallback ke multipart
    if not is_gcs_available():
        return {
            "success": True,
            "use_gcs": False,
            "fallback_to_multipart": True,
            "message": "GCS belum dikonfigurasi pada server ini. Sistem beralih otomatis ke upload langsung."
        }

    signed_info = generate_upload_signed_url(blob_path, content_type=content_type, expiration_minutes=15)
    if not signed_info:
        return {
            "success": True,
            "use_gcs": False,
            "fallback_to_multipart": True,
            "message": "Gagal membuat Signed URL. Sistem beralih ke upload langsung."
        }

    return {
        "success": True,
        "use_gcs": True,
        "upload_url": signed_info["upload_url"],
        "blob_name": signed_info["blob_name"],
        "method": "PUT",
        "expires_in": signed_info["expires_in"],
        "content_type": content_type,
    }


@router.post("/process-job", status_code=status.HTTP_202_ACCEPTED)
async def process_gcs_job(req: ProcessJobRequest):
    """
    Menerima notifikasi setelah upload ke GCS selesai dan memulai pemrosesan asinkronus.
    Memverifikasi integritas Magic Bytes b"%PDF-" dan sanitasi path blob.
    """
    blob_name = req.blob_name.strip()
    action = req.action.lower().strip()
    opts = req.options or {}

    # Validasi Ketat Path Blob (Mencegah Directory Traversal & Akses Objek di Luar uploads/)
    if not re.match(r"^uploads/[a-f0-9]{12}/[\w\-_\. ]+\.pdf$", blob_name, re.IGNORECASE):
        raise HTTPException(
            status_code=400,
            detail="Pola path blob berkas tidak sah. Akses ditolak demi keamanan dokumen."
        )

    base_name = get_safe_base_name(blob_name)
    raw_filename = os.path.basename(blob_name)

    job_id = create_job(message=f"Mempersiapkan pemrosesan berkas {action.upper()} dari penyimpanan awan...")

    async def _run_gcs_task():
        try:
            update_job(job_id, status="processing", progress=5, message="Mengunduh berkas dari Google Cloud Storage...")
            pdf_bytes = await asyncio.to_thread(download_blob_to_bytes, blob_name)
            if not pdf_bytes:
                raise ValueError("Berkas tidak ditemukan atau belum selesai diunggah ke penyimpanan awan.")

            # Validasi Magic Bytes PDF resmi
            if not validate_pdf_bytes(pdf_bytes, raise_exception=False):
                raise ValueError("Berkas yang diunggah tidak valid atau rusak (Header binary PDF tidak dikenali).")

            update_job(job_id, progress=15, message="Memverifikasi integritas dokumen...")

            # ROUTING PEMROSESAN SESUAI AKSI
            if action == "word":
                await _process_word(job_id, pdf_bytes, base_name, opts)
            elif action == "excel":
                await _process_excel(job_id, pdf_bytes, base_name, opts)
            elif action == "ppt":
                await _process_ppt(job_id, pdf_bytes, base_name, opts)
            elif action == "image":
                await _process_image(job_id, pdf_bytes, base_name, opts)
            elif action == "ocr":
                await _process_ocr(job_id, pdf_bytes, base_name, opts)
            elif action == "translate":
                await _process_translate(job_id, pdf_bytes, base_name, opts)
            else:
                raise ValueError(f"Aksi '{action}' tidak didukung.")

        except Exception as e:
            logging.error(f"[GCS Job {job_id}] Kesalahan pemrosesan: {e}")
            update_job(job_id, status="error", error=f"Gagal memproses dokumen: {str(e)}", message="Terjadi kesalahan.")
        finally:
            # Hapus blob sementara di GCS untuk kepatuhan privasi zero-retention
            await asyncio.to_thread(delete_blob, blob_name)

    create_tracked_task(_run_gcs_task())
    return JSONResponse({"job_id": job_id, "status": "pending", "message": f"Tugas {action} dimulai..."}, status_code=202)


@router.post("/configure-security")
async def trigger_gcs_security_config(
    authorization: Optional[str] = Header(None),
    x_admin_email: Optional[str] = Header(None)
):
    """Mengonfigurasi Lifecycle 24-jam dan CORS terproteksi pada bucket GCS (Khusus Admin)."""
    from app.routers.admin import verify_admin_access
    await verify_admin_access(authorization, x_admin_email)
    success = configure_bucket_security()
    return {
        "status": "ok" if success else "warning",
        "message": "Lifecycle policy 24 jam & CORS berhasil diterapkan pada bucket." if success else "Izin service account terbatas. Silakan jalankan script infra/apply-gcs-security.ps1."
    }



# ===========================================================================
# HELPER PEMROSES INTERNAL GCS
# ===========================================================================

async def _process_word(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    tmp_dir = tempfile.mkdtemp()
    docx_filename = f"{base_name}.docx"
    tmp_pdf = os.path.join(tmp_dir, f"input_{base_name}.pdf")
    tmp_docx = os.path.join(tmp_dir, docx_filename)

    def _convert():
        with open(tmp_pdf, "wb") as f:
            f.write(pdf_bytes)

        update_job(job_id, progress=25, message="Menganalisis layout dan tabel dokumen...")
        cv = Converter(tmp_pdf)
        try:
            update_job(job_id, progress=50, message="Mengonversi halaman ke format Word yang dapat diedit...")
            cv.convert(
                tmp_docx,
                connected_border_tolerance=0.5,
                line_overlap_threshold=0.9,
                line_break_free_space_ratio=0.1,
                line_separate_threshold=5.0,
                new_paragraph_free_space_ratio=0.85,
            )
        finally:
            cv.close()

        update_job(job_id, progress=90, message="Merakit berkas Word final...")
        with open(tmp_docx, "rb") as out_f:
            return out_f.read()

    try:
        res = await asyncio.to_thread(_convert)
        update_job(
            job_id,
            status="done",
            progress=100,
            message="Konversi Word selesai!",
            result=res,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=docx_filename,
        )
    finally:
        cleanup_folder(tmp_dir)


async def _process_excel(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    tmp_dir = tempfile.mkdtemp()
    xlsx_filename = f"{base_name}.xlsx"
    tmp_pdf = os.path.join(tmp_dir, f"input_{base_name}.pdf")
    tmp_xlsx = os.path.join(tmp_dir, xlsx_filename)
    mode = opts.get("mode", "tables_only")
    sheet_per_page = opts.get("sheet_per_page", False)

    def _convert():
        with open(tmp_pdf, "wb") as f:
            f.write(pdf_bytes)

        update_job(job_id, progress=25, message="Mendeteksi dan mengekstrak tabel...")
        doc = fitz.open(tmp_pdf)
        try:
            num_pages = len(doc)
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Hasil Ekstraksi"
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )
            header_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
            header_font = Font(bold=True, name='Calibri')
            header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
            data_align = Alignment(vertical='top', wrap_text=True)

            from app.routers.convert import parse_cell_value

            def is_tbl_hdr(row):
                if not row or not any(row):
                    return False
                r_strs = [str(c or "").strip().lower() for c in row if c is not None and str(c).strip()]
                if not r_strs:
                    return False
                if re.match(r'^\d+[\.\)]?$', r_strs[0]):
                    if int(re.sub(r'\D', '', r_strs[0])) > 0:
                        return False
                kws = {'no', 'no.', 'nomor', 'nip', 'nik', 'nama', 'unit kerja', 'jadwal', 'nilai', 'tujuan', 'standar'}
                return sum(1 for s in r_strs if any(k in s for k in kws)) >= 2

            def is_same_hdr(h1, h2):
                if not h1 or not h2 or len(h1) != len(h2):
                    return False
                matches = sum(1 for a, b in zip(h1, h2) if str(a or '').strip().lower() == str(b or '').strip().lower() and str(a or '').strip())
                return matches >= max(2, len(h1) // 2)

            def write_cell(cell, raw, is_hdr=False):
                p = parse_cell_value(raw)
                cell.border = thin_border
                if is_hdr:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_align
                    cell.value = str(p)
                    cell.number_format = '@'
                else:
                    cell.alignment = data_align
                    if isinstance(p, str):
                        cell.number_format = '@'
                        cell.value = p
                    else:
                        cell.value = p

            cur_row = 1
            curr_cols = 0
            curr_hdr = None

            for page_idx in range(num_pages):
                page = doc[page_idx]
                if hasattr(page, 'find_tables'):
                    tabs = page.find_tables()
                    for tab in tabs:
                        extracted = tab.extract()
                        if not extracted or len(extracted) == 0:
                            continue
                        f_row = extracted[0]
                        f_is_hdr = is_tbl_hdr(f_row)
                        is_cont = False
                        if curr_cols > 0 and len(f_row) == curr_cols:
                            if f_is_hdr and curr_hdr and is_same_hdr(curr_hdr, f_row):
                                is_cont = True
                            elif not f_is_hdr:
                                is_cont = True

                        if is_cont:
                            start_r = 1 if f_is_hdr else 0
                            for r_idx in range(start_r, len(extracted)):
                                for c_idx, val in enumerate(extracted[r_idx], 1):
                                    write_cell(ws.cell(row=cur_row, column=c_idx), val, is_hdr=False)
                                cur_row += 1
                        else:
                            if cur_row > 1:
                                cur_row += 2
                            curr_cols = len(f_row)
                            if f_is_hdr:
                                curr_hdr = f_row
                                for c_idx, val in enumerate(f_row, 1):
                                    write_cell(ws.cell(row=cur_row, column=c_idx), val, is_hdr=True)
                                cur_row += 1
                                d_start = 1
                            else:
                                curr_hdr = None
                                d_start = 0
                            for r_idx in range(d_start, len(extracted)):
                                for c_idx, val in enumerate(extracted[r_idx], 1):
                                    write_cell(ws.cell(row=cur_row, column=c_idx), val, is_hdr=False)
                                cur_row += 1

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        first_line = str(cell.value).split('\n')[0]
                        max_len = max(max_len, len(first_line))
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

            wb.save(tmp_xlsx)
        finally:
            if doc and not doc.is_closed:
                doc.close()

        with open(tmp_xlsx, "rb") as out_f:
            return out_f.read()

    try:
        res = await asyncio.to_thread(_convert)
        update_job(
            job_id,
            status="done",
            progress=100,
            message="Konversi Excel selesai!",
            result=res,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=xlsx_filename,
        )
    finally:
        cleanup_folder(tmp_dir)


async def _process_ppt(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    tmp_dir = tempfile.mkdtemp()
    ppt_filename = f"{base_name}.pptx"
    tmp_pdf = os.path.join(tmp_dir, f"input_{base_name}.pdf")
    tmp_ppt = os.path.join(tmp_dir, ppt_filename)

    def _convert():
        with open(tmp_pdf, "wb") as f:
            f.write(pdf_bytes)

        update_job(job_id, progress=30, message="Menyusun slide presentasi PowerPoint...")
        prs = Presentation()
        doc = fitz.open(tmp_pdf)
        clean_doc = fitz.open(tmp_pdf)
        layout_mode = opts.get("layout_mode", "editable")
        try:
            if len(doc) > 0:
                prs.slide_width = int((doc[0].rect.width / 72) * 914400)
                prs.slide_height = int((doc[0].rect.height / 72) * 914400)

            blank = prs.slide_layouts[6]
            for i, page in enumerate(doc):
                slide = prs.slides.add_slide(blank)
                if layout_mode == "visual":
                    pix = page.get_pixmap(dpi=150)
                    img_p = os.path.join(tmp_dir, f"bg_{i}.png")
                    pix.save(img_p)
                    slide.shapes.add_picture(img_p, 0, 0, width=prs.slide_width, height=prs.slide_height)
                else:
                    # 1. Ekstrak latar belakang vektor/grafis jika ada
                    drawings = page.get_drawings()
                    if drawings and len(drawings) > 0 and clean_doc and i < len(clean_doc):
                        try:
                            clean_page = clean_doc[i]
                            for w in clean_page.get_text("words"):
                                clean_page.add_redact_annot(fitz.Rect(w[:4]), fill=False)
                            for img_info in clean_page.get_images():
                                for img_rect in clean_page.get_image_rects(img_info[0]):
                                    clean_page.add_redact_annot(img_rect, fill=False)
                            clean_page.apply_redactions()

                            bg_pix = clean_page.get_pixmap(dpi=180)
                            bg_p = os.path.join(tmp_dir, f"bg_vec_{i}.png")
                            bg_pix.save(bg_p)
                            slide.shapes.add_picture(bg_p, 0, 0, width=prs.slide_width, height=prs.slide_height)
                        except Exception as vec_err:
                            logging.warning(f"Gagal mengekstrak background vektor slide {i+1}: {vec_err}")

                    # 2. Ekstrak gambar bitmap terpisah
                    img_blocks = [b for b in page.get_text("dict", flags=fitz.TEXT_PRESERVE_IMAGES)["blocks"] if b['type'] == 1]
                    for b in img_blocks:
                        img_path = os.path.join(tmp_dir, f"img_{os.urandom(4).hex()}.{b.get('ext', 'png')}")
                        with open(img_path, "wb") as f:
                            f.write(b["image"])
                        x0, y0, x1, y1 = b["bbox"]
                        try:
                            slide.shapes.add_picture(
                                img_path,
                                Inches(x0 / 72),
                                Inches(y0 / 72),
                                width=Inches((x1 - x0) / 72),
                                height=Inches((y1 - y0) / 72)
                            )
                        except Exception:
                            pass

                    # 3. Ekstrak blok teks yang dapat diedit
                    text_blocks = [b for b in page.get_text("dict")["blocks"] if b['type'] == 0]
                    for b in text_blocks:
                        bx0, by0, bx1, by1 = b["bbox"]
                        width = max(bx1 - bx0, 30)
                        height = max(by1 - by0, 15)
                        txBox = slide.shapes.add_textbox(
                            Inches(bx0 / 72),
                            Inches(by0 / 72),
                            Inches(width / 72),
                            Inches(height / 72)
                        )
                        tf = txBox.text_frame
                        tf.word_wrap = True
                        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
                        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

                        first_p = True
                        for line in b["lines"]:
                            p = tf.paragraphs[0] if first_p else tf.add_paragraph()
                            first_p = False
                            for span in line["spans"]:
                                if not span["text"]:
                                    continue
                                run = p.add_run()
                                run.text = span["text"]
                                run.font.size = Pt(max(span["size"], 9))
                                try:
                                    c = span["color"]
                                    run.font.color.rgb = RGBColor((c >> 16) & 0xFF, (c >> 8) & 0xFF, c & 0xFF)
                                except Exception:
                                    pass
                                if span.get("flags", 0) & 16:
                                    run.font.bold = True
                                if span.get("flags", 0) & 2:
                                    run.font.italic = True
                                if "bold" in span.get("font", "").lower():
                                    run.font.bold = True
        finally:
            if clean_doc and not clean_doc.is_closed:
                clean_doc.close()
            if doc and not doc.is_closed:
                doc.close()

        prs.save(tmp_ppt)

        with open(tmp_ppt, "rb") as out_f:
            return out_f.read()

    try:
        res = await asyncio.to_thread(_convert)
        update_job(
            job_id,
            status="done",
            progress=100,
            message="Konversi PowerPoint selesai!",
            result=res,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            filename=ppt_filename,
        )
    finally:
        cleanup_folder(tmp_dir)


async def _process_image(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    fmt = opts.get("output_format", "jpg").lower()
    target_dpi = 300 if opts.get("dpi", 150) >= 300 else 150

    def _convert():
        update_job(job_id, progress=25, message="Merender halaman PDF ke citra resolusi tinggi...")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        try:
            num_pages = len(doc)
            if num_pages == 1:
                pix = doc[0].get_pixmap(dpi=target_dpi)
                b = pix.tobytes("jpeg" if fmt == "jpg" else "png")
                return b, f"image/{'jpeg' if fmt == 'jpg' else 'png'}", f"{base_name}.{fmt}"

            zip_buffer = io.BytesIO()
            with ZipFile(zip_buffer, 'w', compression=ZIP_DEFLATED) as zipf:
                for i, page in enumerate(doc):
                    pct = int(25 + (i / num_pages) * 65)
                    update_job(job_id, progress=pct, message=f"Merender halaman {i + 1} dari {num_pages}...")
                    pix = page.get_pixmap(dpi=target_dpi)
                    zipf.writestr(f"page_{i + 1}.{fmt}", pix.tobytes("jpeg" if fmt == "jpg" else "png"))
            return zip_buffer.getvalue(), "application/zip", f"{base_name}_images.zip"
        finally:
            if doc and not doc.is_closed:
                doc.close()

    res, mtype, fname = await asyncio.to_thread(_convert)
    update_job(job_id, status="done", progress=100, message="Konversi gambar selesai!", result=res, media_type=mtype, filename=fname)


async def _process_ocr(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    clean_langs = opts.get("languages", "ind+eng")
    out_fmt = opts.get("output_format", "pdf")

    def _convert():
        update_job(job_id, progress=20, message=f"Menjalankan OCR ({clean_langs})...")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        searchable_doc = fitz.open()
        try:
            doc_len = len(doc)
            texts = []

            for p_idx in range(doc_len):
                pct = int(20 + (p_idx / doc_len) * 70)
                update_job(job_id, progress=pct, message=f"Menganalisis halaman {p_idx + 1} dari {doc_len}...")
                page = doc[p_idx]
                textpage = page.get_textpage_ocr(language=clean_langs, dpi=150, full=True)
                ocr_text = textpage.extractText()
                texts.append(f"--- Halaman {p_idx + 1} ---\n" + ocr_text.strip())

                new_p = searchable_doc.new_page(width=page.rect.width, height=page.rect.height)
                pix = page.get_pixmap(dpi=150)
                new_p.insert_image(page.rect, stream=pix.tobytes("png"))

                for w in textpage.extractWORDS():
                    w_rect = fitz.Rect(w[0], w[1], w[2], w[3])
                    new_p.insert_text(fitz.Point(w_rect.x0, w_rect.y1), w[4], fontsize=max(6.0, w_rect.height * 0.9), fontname="helv", render_mode=3)

            sample = "\n".join(texts)[:500].strip()

            if out_fmt == "txt":
                return "\n\n".join(texts).encode("utf-8"), "text/plain; charset=utf-8", f"ocr-{base_name}.txt", sample

            out_pdf = searchable_doc.tobytes(garbage=3, deflate=True)
            return out_pdf, "application/pdf", f"searchable-{base_name}.pdf", sample
        finally:
            if doc and not doc.is_closed:
                doc.close()
            if searchable_doc and not searchable_doc.is_closed:
                searchable_doc.close()

    res, mtype, fname, sample = await asyncio.to_thread(_convert)
    update_job(job_id, status="done", progress=100, message="OCR selesai!", result=res, media_type=mtype, filename=fname, sample=sample)


async def _process_translate(job_id: str, pdf_bytes: bytes, base_name: str, opts: dict):
    src_l = opts.get("source_lang", "auto")
    tgt_l = opts.get("target_lang", "id")
    out_fmt = opts.get("output_format", "pdf")

    def _convert():
        update_job(job_id, progress=20, message="Menganalisis tata letak dokumen...")
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        try:
            doc_len = len(doc)
            all_texts = []

            for p_idx in range(doc_len):
                pct = int(20 + (p_idx / doc_len) * 70)
                update_job(job_id, progress=pct, message=f"Menerjemahkan halaman {p_idx + 1} dari {doc_len}...")
                page = doc[p_idx]
                blocks = [b for b in page.get_text("blocks") if len(b) >= 7 and b[6] == 0 and b[4].strip()]
                p_trans = []
                for b in blocks:
                    orig = b[4].strip()
                    if not orig:
                        continue
                    tr = translate_text_chunk(orig, src_l, tgt_l)
                    p_trans.append(tr)
                    if out_fmt == "pdf":
                        rect = fitz.Rect(b[0], b[1], b[2], b[3])
                        page.add_redact_annot(rect, fill=(1, 1, 1))
                        page.apply_redactions()
                        page.insert_textbox(rect, tr, fontsize=9.0, fontname="helv", color=(0, 0, 0), align=0)
                if p_trans:
                    all_texts.append(f"--- Halaman {p_idx + 1} ---\n" + "\n\n".join(p_trans))

            sample = "\n".join(all_texts)[:500].strip()

            if out_fmt == "txt":
                return "\n\n".join(all_texts).encode("utf-8"), "text/plain; charset=utf-8", f"translated-{base_name}.txt", sample

            pdf_out = doc.tobytes(garbage=3, deflate=True)
            return pdf_out, "application/pdf", f"translated-{base_name}.pdf", sample
        finally:
            if doc and not doc.is_closed:
                doc.close()

    res, mtype, fname, sample = await asyncio.to_thread(_convert)
    update_job(job_id, status="done", progress=100, message="Terjemahan selesai!", result=res, media_type=mtype, filename=fname, sample=sample)
